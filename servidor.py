from flask import Flask, request, jsonify, render_template_string
from datetime import datetime
import time
import os
from openpyxl import Workbook, load_workbook

app = Flask(__name__)

# ==============================================================================
# ARCHIVO EXCEL DE REGISTRO
# ==============================================================================

ARCHIVO_EXCEL = "registro_nfc.xlsx"

ENCABEZADOS = [
    "ID", "Fecha/Hora", "UID", "Tipo",
    "Operacion", "Detalle", "Autorizado", "Latencia (ms)"
]


def inicializar_excel():
    """Crea el archivo Excel con encabezados si no existe todavia."""
    if not os.path.exists(ARCHIVO_EXCEL):
        wb = Workbook()
        ws = wb.active
        ws.title = "Registro NFC"
        ws.append(ENCABEZADOS)
        wb.save(ARCHIVO_EXCEL)


def guardar_evento_excel(evento):
    """Agrega una fila nueva al archivo Excel con el evento recibido."""
    try:
        wb = load_workbook(ARCHIVO_EXCEL)
        ws = wb.active
        ws.append([
            evento["id"],
            evento["fecha"],
            evento["uid"],
            evento["tipo"],
            evento["operacion"],
            evento["detalle"],
            evento["autorizado"],
            evento["latencia"]
        ])
        wb.save(ARCHIVO_EXCEL)
    except Exception as e:
        print(f"ERROR guardando en Excel: {e}")


# ==============================================================================
# BASE DE DATOS Y ESTADO EN MEMORIA
# ==============================================================================

# Ahora inicia vacia, sin usuarios de ejemplo
usuarios_db = {}

productos_db = {
    "1": {"nombre": "Cola", "precio": 2.50, "stock": 10},
    "2": {"nombre": "Pepsi", "precio": 1.00, "stock": 5}
}

logs_pruebas = []

estado_sistema = {
    "ultimo_usuario": "Ninguno",
    "ultima_operacion": "Ninguna"
}


def registrar_evento(uid, operacion, detalle, autorizado, latencia_ms, tipo_tag="Mifare Classic"):
    estado_sistema["ultimo_usuario"] = usuarios_db.get(
        uid, {}).get("nombre", uid)
    estado_sistema["ultima_operacion"] = f"{operacion}: {detalle}"

    evento = {
        "id": len(logs_pruebas) + 1,
        "fecha": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "uid": uid,
        "tipo": tipo_tag,
        "operacion": operacion,
        "detalle": detalle,
        "autorizado": "Sí" if autorizado else "No",
        "latencia": round(latencia_ms, 2)
    }

    logs_pruebas.append(evento)

    # Guardar tambien en el archivo Excel
    guardar_evento_excel(evento)

# ==============================================================================
# INTERFAZ WEB CON ACTUALIZACIÓN EN TIEMPO REAL
# ==============================================================================


HTML_TEMPLATE = """
<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Servidor NFC - Máquina de Cobro</title>
    <style>
        body { font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; margin: 30px; background-color: #f8fafc; color: #0f172a; }
        h1 { font-size: 26px; font-weight: 700; margin-bottom: 20px; }
        h2 { font-size: 18px; font-weight: 600; margin-top: 30px; margin-bottom: 10px; color: #1e293b; }
        .card { background: #e2e8f0; padding: 18px; border-radius: 8px; margin-bottom: 20px; }
        .card p { margin: 6px 0; font-size: 15px; }
        table { width: 100%; border-collapse: collapse; background: #ffffff; border-radius: 6px; overflow: hidden; margin-top: 8px; box-shadow: 0 1px 3px rgba(0,0,0,0.05); }
        th, td { text-align: center; padding: 10px 14px; border: 1px solid #cbd5e1; font-size: 13px; }
        th { background-color: #f1f5f9; font-weight: 600; color: #334155; }
        tr:nth-child(even) { background-color: #f8fafc; }
        .status-yes { color: #16a34a; font-weight: bold; }
        .status-no { color: #dc2626; font-weight: bold; }
        .btn-edit { cursor: pointer; border: none; background: none; font-size: 14px; margin-left: 8px; }
        .btn-edit:hover { transform: scale(1.2); }
    </style>
</head>
<body>
    <h1>Sistema NFC - Máquina de cobro</h1>

    <div class="card">
        <h2>Estado del sistema</h2>
        <p><strong>Usuario activo:</strong> <span id="estado-usuario">-</span></p>
        <p><strong>Última operación:</strong> <span id="estado-operacion">-</span></p>
    </div>

    <h2>Usuarios registrados</h2>
    <table>
        <thead>
            <tr>
                <th>Usuario</th>
                <th>UID NFC</th>
                <th>Tipo</th>
                <th>Saldo</th>
            </tr>
        </thead>
        <tbody id="tbody-usuarios"></tbody>
    </table>

    <h2>Productos</h2>
    <table>
        <thead>
            <tr>
                <th>Código (Teclado)</th>
                <th>Producto</th>
                <th>Precio</th>
                <th>Stock</th>
            </tr>
        </thead>
        <tbody id="tbody-productos"></tbody>
    </table>

    <h2> Registro de pruebas de lectura RFID/NFC</h2>
    <table>
        <thead>
            <tr>
                <th>Nº</th>
                <th>Fecha / Hora</th>
                <th>UID / Token</th>
                <th>Tipo</th>
                <th>Operación / Detalle</th>
                <th>Autorizado</th>
                <th>Latencia (ms)</th>
            </tr>
        </thead>
        <tbody id="tbody-logs"></tbody>
    </table>

    <script>
        function cambiarNombre(uid) {
            let nuevoNombre = prompt("Ingresa el nuevo nombre para este usuario:");
            if (nuevoNombre && nuevoNombre.trim() !== "") {
                fetch('/actualizar_nombre', {
                    method: 'POST',
                    headers: { 'Content-Type': 'application/json' },
                    body: JSON.stringify({ uid: uid, nombre: nuevoNombre.trim() })
                })
                .then(response => response.json())
                .then(data => cargarDatos()); 
            }
        }

        async function cargarDatos() {
            try {
                const res = await fetch('/api/datos');
                const data = await res.json();

                document.getElementById('estado-usuario').innerText = data.estado.ultimo_usuario;
                document.getElementById('estado-operacion').innerText = data.estado.ultima_operacion;

                let trUsuarios = '';
                for (const [uid, user] of Object.entries(data.usuarios)) {
                    trUsuarios += `<tr>
                        <td>${user.nombre} <button class="btn-edit" title="Editar Nombre" onclick="cambiarNombre('${uid}')">✏️</button></td>
                        <td>${uid}</td>
                        <td>${user.tipo}</td>
                        <td>$${user.saldo.toFixed(2)}</td>
                    </tr>`;
                }
                document.getElementById('tbody-usuarios').innerHTML = trUsuarios;

                let trProductos = '';
                for (const [codigo, prod] of Object.entries(data.productos)) {
                    trProductos += `<tr>
                        <td>${codigo}</td>
                        <td>${prod.nombre}</td>
                        <td>$${prod.precio.toFixed(2)}</td>
                        <td>${prod.stock}</td>
                    </tr>`;
                }
                document.getElementById('tbody-productos').innerHTML = trProductos;

                let trLogs = '';
                data.logs.forEach(log => {
                    let statusClass = log.autorizado === 'Sí' ? 'status-yes' : 'status-no';
                    trLogs += `<tr>
                        <td>${log.id}</td>
                        <td>${log.fecha}</td>
                        <td>${log.uid}</td>
                        <td>${log.tipo}</td>
                        <td><strong>${log.operacion}:</strong> ${log.detalle}</td>
                        <td class="${statusClass}">${log.autorizado}</td>
                        <td>${log.latencia} ms</td>
                    </tr>`;
                });
                document.getElementById('tbody-logs').innerHTML = trLogs;

            } catch (error) {
                console.error("Error obteniendo datos del servidor:", error);
            }
        }

        window.onload = cargarDatos;
        setInterval(cargarDatos, 1500);
    </script>
</body>
</html>
"""


@app.route('/', methods=['GET'])
def inicio():
    return render_template_string(HTML_TEMPLATE)


@app.route('/api/datos', methods=['GET'])
def api_datos():
    return jsonify({
        "estado": estado_sistema,
        "usuarios": usuarios_db,
        "productos": productos_db,
        "logs": logs_pruebas
    })


@app.route('/actualizar_nombre', methods=['POST'])
def actualizar_nombre():
    data = request.get_json() or {}
    uid = str(data.get('uid', '')).strip().upper()
    nuevo_nombre = str(data.get('nombre', '')).strip()

    if uid in usuarios_db and nuevo_nombre:
        usuarios_db[uid]["nombre"] = nuevo_nombre
        if estado_sistema["ultimo_usuario"] == uid or estado_sistema["ultimo_usuario"] != "Ninguno":
            estado_sistema["ultimo_usuario"] = nuevo_nombre
        return jsonify({"mensaje": "Nombre actualizado"}), 200

    return jsonify({"error": "No se pudo actualizar"}), 400


@app.route('/comprar', methods=['POST'])
def comprar_producto():
    t_start = time.time()
    data = request.get_json() or {}
    uid = str(data.get('uid', '')).strip().upper()
    producto_key = str(data.get('producto', ''))
    tipo_tag = data.get('tipo', 'Mifare Classic')

    if not uid or uid not in usuarios_db:
        lat = (time.time() - t_start) * 1000
        registrar_evento(uid, "COMPRA", "Tarjeta no registrada",
                         False, lat, tipo_tag)
        return jsonify({"error": "Usuario no registrado"}), 404

    if producto_key not in productos_db:
        lat = (time.time() - t_start) * 1000
        registrar_evento(
            uid, "COMPRA", f"Opción inválida ({producto_key})", False, lat, tipo_tag)
        return jsonify({"error": "Producto invalido"}), 400

    usuario = usuarios_db[uid]
    producto = productos_db[producto_key]

    if producto["stock"] <= 0 or usuario["saldo"] < producto["precio"]:
        lat = (time.time() - t_start) * 1000
        motivo = "Sin Stock" if producto["stock"] <= 0 else "Saldo Insuficiente"
        registrar_evento(
            uid, "COMPRA", f"Rechazado ({motivo})", False, lat, tipo_tag)
        return jsonify({"error": motivo, "saldo": round(usuario["saldo"], 2)}), 400

    usuario["saldo"] -= producto["precio"]
    producto["stock"] -= 1
    lat = (time.time() - t_start) * 1000

    registrar_evento(
        uid, "COMPRA", f"Compró {producto['nombre']} por ${producto['precio']:.2f}", True, lat, tipo_tag)

    return jsonify({
        "mensaje": "Compra exitosa",
        "producto": producto["nombre"],
        "saldo": round(usuario["saldo"], 2)
    }), 200


@app.route('/recargar', methods=['POST'])
def recargar_saldo():
    t_start = time.time()
    data = request.get_json() or {}
    uid = str(data.get('uid', '')).strip().upper()
    monto = float(data.get('monto', 0))
    tipo_tag = data.get('tipo', 'Mifare Classic')

    if not uid or uid not in usuarios_db:
        lat = (time.time() - t_start) * 1000
        registrar_evento(uid, "RECARGA", "Usuario no existe",
                         False, lat, tipo_tag)
        return jsonify({"error": "Usuario no existe"}), 404

    if monto <= 0:
        lat = (time.time() - t_start) * 1000
        registrar_evento(uid, "RECARGA", "Monto inválido",
                         False, lat, tipo_tag)
        return jsonify({"error": "Monto invalido"}), 400

    usuarios_db[uid]["saldo"] += monto
    lat = (time.time() - t_start) * 1000

    registrar_evento(
        uid, "RECARGA", f"Recargó ${monto:.2f}", True, lat, tipo_tag)

    return jsonify({
        "mensaje": "Recarga exitosa",
        "saldo": round(usuarios_db[uid]["saldo"], 2)
    }), 200


@app.route('/usuario/<path:uid>', methods=['GET'])
def consultar_usuario(uid):
    t_start = time.time()
    uid = uid.strip().upper()

    if uid in usuarios_db:
        usuario = usuarios_db[uid]
        lat = (time.time() - t_start) * 1000
        registrar_evento(
            uid, "CONSULTA", f"Saldo: ${usuario['saldo']:.2f}", True, lat, usuario["tipo"])
        return jsonify({
            "nombre": usuario["nombre"],
            "saldo": round(usuario["saldo"], 2)
        }), 200
    else:
        lat = (time.time() - t_start) * 1000
        registrar_evento(uid, "CONSULTA", "Usuario no encontrado", False, lat)
        return jsonify({"error": "Usuario no registrado"}), 404


@app.route('/registrar', methods=['POST'])
def registrar_usuario():
    t_start = time.time()
    data = request.get_json() or {}
    uid = str(data.get('uid', '')).strip().upper()
    tipo_tag = data.get('tipo', 'Mifare Classic')

    if not uid:
        return jsonify({"error": "UID faltante"}), 400

    if uid in usuarios_db:
        lat = (time.time() - t_start) * 1000
        registrar_evento(
            uid, "REGISTRO", "Tarjeta ya registrada", False, lat, tipo_tag)
        return jsonify({"mensaje": "Usuario ya registrado"}), 409

    num_user = len(usuarios_db) + 1
    nombre_user = f"USER{num_user}"

    usuarios_db[uid] = {
        "nombre": nombre_user,
        "saldo": 0.00,
        "tipo": tipo_tag
    }

    lat = (time.time() - t_start) * 1000
    registrar_evento(
        uid, "REGISTRO", f"Usuario {nombre_user} creado", True, lat, tipo_tag)

    return jsonify({
        "mensaje": "Usuario creado correctamente",
        "uid": uid
    }), 200


if __name__ == '__main__':
    inicializar_excel()
    app.run(host='0.0.0.0', port=5000, debug=True)
